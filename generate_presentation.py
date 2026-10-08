import sys
import os
import base64
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Modern Minimalist High-End
    BG_COLOR = RGBColor(248, 250, 252)        # #F8FAFC (Slate 50)
    CARD_BG = RGBColor(255, 255, 255)         # #FFFFFF
    BORDER_COLOR = RGBColor(226, 232, 240)    # #E2E8F0 (Slate 200)
    TEXT_MAIN = RGBColor(15, 23, 42)          # #0F172A (Slate 900)
    TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B (Slate 500)
    TEXT_BODY = RGBColor(51, 65, 85)          # #334155 (Slate 700)
    ACCENT_BLUE = RGBColor(37, 99, 235)       # #2563EB (Blue 600)
    ACCENT_BG = RGBColor(239, 246, 255)       # #EFF6FF (Blue 50)
    ACCENT_DARK = RGBColor(30, 58, 138)       # #1E3A8A (Blue 900)
    SUCCESS_TEXT = RGBColor(13, 148, 136)     # #0D9488 (Teal 600)

    FONT_FAMILY = "Arial"

    def apply_slide_base(slide, slide_num, total_slides=10):
        # Background rect
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        
        # Footer dividing line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.color.rgb = BORDER_COLOR

        # Footer Left text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(8), Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "MadadgaarAI  |  B.Tech CSE Major Project  |  Hybrid Funding Intelligence System"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED

        # Footer Right text (Slide number)
        tx_num = slide.shapes.add_textbox(Inches(11.5), Inches(6.95), Inches(1.033), Inches(0.35))
        tf_num = tx_num.text_frame
        tf_num.word_wrap = True
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        pn = tf_num.paragraphs[0]
        pn.alignment = PP_ALIGN.RIGHT
        pn.text = f"{str(slide_num).zfill(2)} / {str(total_slides).zfill(2)}"
        pn.font.name = FONT_FAMILY
        pn.font.size = Pt(9.5)
        pn.font.bold = True
        pn.font.color.rgb = TEXT_MUTED

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Tag pill
        tag_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(3.2), Inches(0.32))
        tag_shape.fill.solid()
        tag_shape.fill.fore_color.rgb = ACCENT_BG
        tag_shape.line.color.rgb = BORDER_COLOR
        tf_tag = tag_shape.text_frame
        tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.alignment = PP_ALIGN.CENTER
        p_tag.text = tag_text.upper()
        p_tag.font.name = FONT_FAMILY
        p_tag.font.size = Pt(9)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_BLUE

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.84), Inches(11.733), Inches(0.55))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(21)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        # Subtitle
        if subtitle_text:
            s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.35))
            tf_s = s_box.text_frame
            tf_s.word_wrap = True
            tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
            ps = tf_s.paragraphs[0]
            ps.text = subtitle_text
            ps.font.name = FONT_FAMILY
            ps.font.size = Pt(10.5)
            ps.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title="", title_accent="", fill_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        if title:
            t_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), width - Inches(0.5), Inches(0.4))
            tf = t_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = FONT_FAMILY
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = TEXT_MAIN
            if title_accent:
                p_acc = tf.add_paragraph()
                p_acc.text = title_accent
                p_acc.font.name = FONT_FAMILY
                p_acc.font.size = Pt(9.5)
                p_acc.font.color.rgb = ACCENT_BLUE

        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide1, 1)

    hero = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.65), Inches(11.733), Inches(6.0))
    hero.fill.solid()
    hero.fill.fore_color.rgb = CARD_BG
    hero.line.color.rgb = BORDER_COLOR
    hero.line.width = Pt(1.2)

    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(1.05), Inches(4.2), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT_BG
    badge.line.color.rgb = BORDER_COLOR
    tb = badge.text_frame
    tb.vertical_anchor = MSO_ANCHOR.MIDDLE
    tb.margin_left = tb.margin_top = tb.margin_right = tb.margin_bottom = 0
    pb = tb.paragraphs[0]
    pb.alignment = PP_ALIGN.CENTER
    pb.text = "MAJOR PROJECT PRESENTATION • B.TECH CSE"
    pb.font.name = FONT_FAMILY
    pb.font.size = Pt(9.5)
    pb.font.bold = True
    pb.font.color.rgb = ACCENT_BLUE

    t_box = slide1.shapes.add_textbox(Inches(1.3), Inches(1.5), Inches(10.7), Inches(0.85))
    tf = t_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "MadadgaarAI"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

    s_box = slide1.shapes.add_textbox(Inches(1.3), Inches(2.4), Inches(10.7), Inches(0.8))
    tfs = s_box.text_frame
    tfs.word_wrap = True
    tfs.margin_left = tfs.margin_top = tfs.margin_right = tfs.margin_bottom = 0
    ps = tfs.paragraphs[0]
    ps.text = "Hybrid Retrieval & Autonomous Agent Framework for Student Welfare Schemes"
    ps.font.name = FONT_FAMILY
    ps.font.size = Pt(16)
    ps.font.bold = True
    ps.font.color.rgb = ACCENT_BLUE

    ps2 = tfs.add_paragraph()
    ps2.space_before = Pt(4)
    ps2.text = "Deterministic Eligibility Verification • Multilingual Saral Advisory • Adaptive Document Ingestion"
    ps2.font.name = FONT_FAMILY
    ps2.font.size = Pt(11)
    ps2.font.color.rgb = TEXT_MUTED

    # Bottom Two Equal Cards: Left = Team & Mentor, Right = Core Innovations
    y_cards = Inches(3.45)
    h_card_1 = Inches(2.85)
    w_card_1 = Inches(5.15)
    gap_card_1 = Inches(0.4)

    # Card 1: Team & Mentor
    c_team = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), y_cards, w_card_1, h_card_1)
    c_team.fill.solid()
    c_team.fill.fore_color.rgb = RGBColor(248, 250, 252)
    c_team.line.color.rgb = BORDER_COLOR

    tx_team = slide1.shapes.add_textbox(Inches(1.55), y_cards + Inches(0.2), w_card_1 - Inches(0.5), h_card_1 - Inches(0.4))
    tf_team = tx_team.text_frame
    tf_team.word_wrap = True
    tf_team.margin_left = tf_team.margin_top = tf_team.margin_right = tf_team.margin_bottom = 0
    
    pt0 = tf_team.paragraphs[0]
    pt0.text = "PROJECT TEAM & SUPERVISION"
    pt0.font.name = FONT_FAMILY
    pt0.font.size = Pt(10.5)
    pt0.font.bold = True
    pt0.font.color.rgb = ACCENT_DARK

    pt_sub1 = tf_team.add_paragraph()
    pt_sub1.space_before = Pt(8)
    pt_sub1.text = "Project Members:"
    pt_sub1.font.name = FONT_FAMILY
    pt_sub1.font.size = Pt(9.5)
    pt_sub1.font.bold = True
    pt_sub1.font.color.rgb = TEXT_MUTED

    members = [
        ("Aman Singh", "2023BTECH006"),
        ("Anjisht Amritanshu", "2023BTECH009"),
        ("Ayush Choudhary", "2023BTECH020")
    ]
    for m_name, m_roll in members:
        pm = tf_team.add_paragraph()
        pm.space_before = Pt(3)
        pm.text = f"• {m_name} "
        pm.font.name = FONT_FAMILY
        pm.font.size = Pt(10)
        pm.font.bold = True
        pm.font.color.rgb = TEXT_MAIN
        rm = pm.add_run()
        rm.text = f"({m_roll})"
        rm.font.name = FONT_FAMILY
        rm.font.size = Pt(9.5)
        rm.font.bold = False
        rm.font.color.rgb = TEXT_MUTED

    pt_mentor = tf_team.add_paragraph()
    pt_mentor.space_before = Pt(10)
    pt_mentor.text = "Project Mentor: "
    pt_mentor.font.name = FONT_FAMILY
    pt_mentor.font.size = Pt(10)
    pt_mentor.font.bold = True
    pt_mentor.font.color.rgb = ACCENT_BLUE

    rm_m = pt_mentor.add_run()
    rm_m.text = "Prof. Dr. Deepika Prakash"
    rm_m.font.name = FONT_FAMILY
    rm_m.font.size = Pt(10.5)
    rm_m.font.bold = True
    rm_m.font.color.rgb = TEXT_MAIN

    # Card 2: Core Engineering Pillars
    x_card_r = Inches(1.3) + w_card_1 + gap_card_1
    c_arch = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_card_r, y_cards, w_card_1, h_card_1)
    c_arch.fill.solid()
    c_arch.fill.fore_color.rgb = RGBColor(248, 250, 252)
    c_arch.line.color.rgb = BORDER_COLOR

    tx_arch = slide1.shapes.add_textbox(x_card_r + Inches(0.25), y_cards + Inches(0.2), w_card_1 - Inches(0.5), h_card_1 - Inches(0.4))
    tf_arch = tx_arch.text_frame
    tf_arch.word_wrap = True
    tf_arch.margin_left = tf_arch.margin_top = tf_arch.margin_right = tf_arch.margin_bottom = 0

    pa0 = tf_arch.paragraphs[0]
    pa0.text = "CORE ENGINEERING INNOVATIONS"
    pa0.font.name = FONT_FAMILY
    pa0.font.size = Pt(10.5)
    pa0.font.bold = True
    pa0.font.color.rgb = ACCENT_DARK

    pillars = [
        ("01 | Dual-Retriever Search", "BM25Okapi lexical precision + 384-d dense vectors (all-MiniLM-L6-v2) fused via Reciprocal Rank Fusion (k=60)."),
        ("02 | Zero-Hallucination Gate", "Strict deterministic rules evaluating marks cutoff, income ceiling, and quotas before recommendation."),
        ("03 | Saral Advisory & OCR", "Plain Hinglish guidance, automated document checklists, and adaptive density OCR routing (rho < 120).")
    ]
    for p_title, p_desc in pillars:
        pa_t = tf_arch.add_paragraph()
        pa_t.space_before = Pt(8)
        pa_t.text = f"• {p_title}"
        pa_t.font.name = FONT_FAMILY
        pa_t.font.size = Pt(10)
        pa_t.font.bold = True
        pa_t.font.color.rgb = TEXT_MAIN

        pa_d = tf_arch.add_paragraph()
        pa_d.space_before = Pt(2)
        pa_d.text = f"  {p_desc}"
        pa_d.font.name = FONT_FAMILY
        pa_d.font.size = Pt(9)
        pa_d.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide2, 2)
    add_header(slide2, "Background & Motivation", "The Scholarship Discovery & Distribution Gap", 
               "Thousands of eligible students fail to secure financial aid due to severe systemic bottlenecks in scheme discovery.")

    col_w_2 = Inches(3.68)
    gap_2 = Inches(0.34)
    y_pos_2 = Inches(1.9)
    card_h_2 = Inches(4.75)

    problems = [
        ("01. Information Fragmentation", "Disparate Portals & Schemes", [
             ("2,500+ Active Schemes", "Distributed across National Scholarship Portal (NSP), state welfare departments, and corporate CSRs with no unified index."),
             ("Inconsistent Formats", "Announcements published across PDF gazettes, regional press releases, and isolated department pages."),
             ("Missed Deadlines", "Students miss application windows due to scattered notifications and lack of proactive deadline alerts.")
         ]),
        ("02. High Bureaucratic Friction", "Strict Multi-Clause Eligibility", [
             ("Complex Eligibility Rules", "Intersecting criteria across academic percentage, family annual income cap, domicile, caste quota, and gender."),
             ("High Application Rejection", "Students spend weeks collecting certificates only to be rejected over ambiguous sub-clauses."),
             ("Fund Underutilization", "Crores in government and philanthropic welfare budgets lapse unallocated each academic cycle.")
         ]),
        ("03. Linguistic & Digital Divide", "Legalistic Phrasing & Paperwork", [
             ("Complex Legal English/Hindi", "Official circulars written in dense bureaucratic terminology alienating Tier-2/3 & rural students."),
             ("Lack of Document Guidance", "Ambiguity on which specific documents (Tehsildar income certificate vs. salary slip) are mandatory."),
             ("Exploitative Cyber Cafes", "Students depend on fee-charging middlemen and cyber cafes to decipher circulars and apply.")
         ])
    ]

    for i, (p_tag, p_head, bullets) in enumerate(problems):
        x = Inches(0.8) + i * (col_w_2 + gap_2)
        add_card(slide2, x, y_pos_2, col_w_2, card_h_2)

        tx = slide2.shapes.add_textbox(x + Inches(0.28), y_pos_2 + Inches(0.25), col_w_2 - Inches(0.56), card_h_2 - Inches(0.5))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = p_tag.upper()
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p_h = tf.add_paragraph()
        p_h.space_before = Pt(4)
        p_h.text = p_head
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_MAIN

        for b_title, b_desc in bullets:
            pb_t = tf.add_paragraph()
            pb_t.space_before = Pt(10)
            pb_t.text = f"• {b_title}"
            pb_t.font.name = FONT_FAMILY
            pb_t.font.size = Pt(10.5)
            pb_t.font.bold = True
            pb_t.font.color.rgb = TEXT_MAIN

            pb_d = tf.add_paragraph()
            pb_d.space_before = Pt(2)
            pb_d.text = f"  {b_desc}"
            pb_d.font.name = FONT_FAMILY
            pb_d.font.size = Pt(9.5)
            pb_d.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 3: PROPOSED SOLUTION & VALUE PROPOSITION
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide3, 3)
    add_header(slide3, "Proposed Solution", "MadadgaarAI: AI-Driven Funding Intelligence Platform",
               "A unified, deterministic intelligence layer connecting students with verified welfare opportunities.")

    w_left_3 = Inches(6.8)
    w_right_3 = Inches(4.6)
    gap_3 = Inches(0.33)
    y_pos_3 = Inches(1.9)
    h_3 = Inches(3.6)

    add_card(slide3, Inches(0.8), y_pos_3, w_left_3, h_3, "Four Core Architectural Pillars", "Engineered for absolute accuracy and zero ambiguity")
    tx_l3 = slide3.shapes.add_textbox(Inches(1.05), y_pos_3 + Inches(0.72), w_left_3 - Inches(0.5), h_3 - Inches(0.8))
    tf_l3 = tx_l3.text_frame
    tf_l3.word_wrap = True
    tf_l3.margin_left = tf_l3.margin_top = tf_l3.margin_right = tf_l3.margin_bottom = 0

    pillars = [
        ("Automated Document Ingestion & OCR", "Extracts scheme guidelines from digital and scanned PDFs using PyPDF with dynamic Tesseract OCR routing."),
        ("Deterministic Eligibility Engine", "Strict algorithmic validation on marks, income cap, and quota restrictions; eliminates LLM hallucination."),
        ("Hybrid Retrieval Engine (Dense + Sparse)", "Combines BM25Okapi lexical precision with all-MiniLM-L6-v2 vector embeddings via Reciprocal Rank Fusion (k=60)."),
        ("Multilingual Saral Advisory & Vidyarthi AI", "Delivers simplified Hinglish eligibility translations, dynamic document checklists, and interactive guidance.")
    ]

    for i, (pil_title, pil_desc) in enumerate(pillars):
        p = tf_l3.paragraphs[0] if i == 0 else tf_l3.add_paragraph()
        if i > 0:
            p.space_before = Pt(8)
        p.text = f"0{i+1}. {pil_title}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        p_sub = tf_l3.add_paragraph()
        p_sub.space_before = Pt(1)
        p_sub.text = f"     {pil_desc}"
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = TEXT_BODY

    add_card(slide3, Inches(0.8) + w_left_3 + gap_3, y_pos_3, w_right_3, h_3, "System Comparison", "Conventional Portals vs. MadadgaarAI")
    tx_r3 = slide3.shapes.add_textbox(Inches(0.8) + w_left_3 + gap_3 + Inches(0.25), y_pos_3 + Inches(0.72), w_right_3 - Inches(0.5), h_3 - Inches(0.8))
    tf_r3 = tx_r3.text_frame
    tf_r3.word_wrap = True
    tf_r3.margin_left = tf_r3.margin_top = tf_r3.margin_right = tf_r3.margin_bottom = 0

    comparisons = [
        ("Portal Discovery", "Static keyword forms", "Hybrid Semantic + BM25 Search"),
        ("Eligibility Guarantee", "Self-checked by user", "Deterministic Rule Engine (0% Hallucination)"),
        ("Notice Readability", "Complex legalistic English/Hindi", "Instant Hinglish Saral Guide summary"),
        ("Document Readiness", "Vague document mentions", "Dynamic step-by-step checklist generation")
    ]

    for i, (dim, conv, mad) in enumerate(comparisons):
        p = tf_r3.paragraphs[0] if i == 0 else tf_r3.add_paragraph()
        if i > 0:
            p.space_before = Pt(7)
        p.text = f"• {dim}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        p_c = tf_r3.add_paragraph()
        p_c.space_before = Pt(1)
        p_c.text = f"  Old: {conv}  ➔  MadadgaarAI: {mad}"
        p_c.font.name = FONT_FAMILY
        p_c.font.size = Pt(9)
        p_c.font.color.rgb = ACCENT_BLUE

    # Bottom Metric Ribbon
    y_bot_3 = Inches(5.65)
    h_bot_3 = Inches(1.05)
    w_stat_3 = Inches(3.68)
    gap_stat_3 = Inches(0.34)

    stats = [
        ("0% Hallucination", "Strict deterministic filters guarantee 100% eligibility accuracy"),
        ("< 150ms Query Latency", "Asynchronous hybrid retrieval over vector & lexical indices"),
        ("Dual-Layer OCR Routing", "Density-based routing ensures zero OCR slowdown on digital PDFs")
    ]

    for i, (s_val, s_lbl) in enumerate(stats):
        x_s = Inches(0.8) + i * (w_stat_3 + gap_stat_3)
        s_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_s, y_bot_3, w_stat_3, h_bot_3)
        s_card.fill.solid()
        s_card.fill.fore_color.rgb = ACCENT_BG
        s_card.line.color.rgb = BORDER_COLOR

        tx_s = slide3.shapes.add_textbox(x_s + Inches(0.2), y_bot_3 + Inches(0.12), w_stat_3 - Inches(0.4), h_bot_3 - Inches(0.24))
        tf_s = tx_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_sv = tf_s.paragraphs[0]
        p_sv.text = s_val
        p_sv.font.name = FONT_FAMILY
        p_sv.font.size = Pt(13)
        p_sv.font.bold = True
        p_sv.font.color.rgb = ACCENT_DARK

        p_sl = tf_s.add_paragraph()
        p_sl.space_before = Pt(2)
        p_sl.text = s_lbl
        p_sl.font.name = FONT_FAMILY
        p_sl.font.size = Pt(9)
        p_sl.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 4: FIVE-LAYER SYSTEM ARCHITECTURE (WITH FIGURE 1)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide4, 4)
    add_header(slide4, "System Architecture", "Five-Layer Modular System Architecture",
               "Decoupled multi-tier design separating external ingestion, persistent storage, matching services, and API gateway.")

    w_fig1_card = Inches(5.8)
    h_fig1_card = Inches(4.85)
    y_fig1 = Inches(1.85)

    add_card(slide4, Inches(0.8), y_fig1, w_fig1_card, h_fig1_card, "System Architecture Model", "Figure 1: Five-Layer Architecture")
    
    fig1_path = "presentation_assets/figure1_system_architecture.png"
    if os.path.exists(fig1_path):
        slide4.shapes.add_picture(fig1_path, Inches(1.05), y_fig1 + Inches(0.68), width=Inches(5.3))

    w_right_4 = Inches(5.6)
    x_right_4 = Inches(6.93)
    add_card(slide4, x_right_4, y_fig1, w_right_4, h_fig1_card, "Five Core Functional Layers", "Technical breakdown from ingestion to delivery")

    tx_4r = slide4.shapes.add_textbox(x_right_4 + Inches(0.25), y_fig1 + Inches(0.7), w_right_4 - Inches(0.5), h_fig1_card - Inches(0.8))
    tf_4r = tx_4r.text_frame
    tf_4r.word_wrap = True
    tf_4r.margin_left = tf_4r.margin_top = tf_4r.margin_right = tf_4r.margin_bottom = 0

    layers = [
        ("Layer 1: External Sources", "NSP Central schemes, State welfare portals (MahaDBT, UP), AICTE/UGC guidelines, DST/CSIR grants, and Corporate CSR trusts."),
        ("Layer 2: Ingestion, Parsing & Normalisation", "Asynchronous crawlers (aiohttp, backoff retries), SHA-256 deduplication, native PDF parser with density check rho, and OCR fallback extractor."),
        ("Layer 3: Storage & Dual Indexing", "SQLite with Write-Ahead Logging (WAL) + JSON/CSV records, BM25Okapi inverted index, and 384-dimensional dense vector embeddings."),
        ("Layer 4: Matching & Decision Services", "Reciprocal Rank Fusion (RRF) hybrid search, Vidyarthi AI student matcher, Saral Hinglish guide, checklist generator, and RFC 5545 calendar export."),
        ("Layer 5: Presentation & API Gateway", "Asynchronous FastAPI REST endpoints (Uvicorn, Pydantic v2 schemas) powering the responsive Neo-Brutalist student dashboard.")
    ]

    for i, (l_title, l_desc) in enumerate(layers):
        p = tf_4r.paragraphs[0] if i == 0 else tf_4r.add_paragraph()
        if i > 0:
            p.space_before = Pt(7)
        p.text = f"• {l_title}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE if i in [1, 2, 3] else TEXT_MAIN

        pd = tf_4r.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {l_desc}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 5: END-TO-END DATA PROCESSING PIPELINE (WITH FIGURE 2)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide5, 5)
    add_header(slide5, "Data Processing Pipeline", "End-to-End Data Processing Pipeline Walkthrough",
               "Step-by-step lifecycle from asynchronous crawling and OCR fallback to dual-index query serving.")

    # Left Container Card for Figure 2
    w_card_5l = Inches(8.1)
    h_card_5 = Inches(4.85)
    y_pos_5 = Inches(1.85)

    add_card(slide5, Inches(0.8), y_pos_5, w_card_5l, h_card_5, "Pipeline Data Flow Diagram", "Figure 2: Execution Flowchart")

    fig2_path = "presentation_assets/figure2_pipeline_flow.png"
    if os.path.exists(fig2_path):
        # Embed Figure 2 image neatly with zero overlap
        slide5.shapes.add_picture(fig2_path, Inches(1.0), y_pos_5 + Inches(0.65), width=Inches(7.7))

    # Bottom 3 micro step summary pills inside the left card
    y_pills = Inches(5.8)
    h_pill = Inches(0.7)
    w_pill = Inches(2.35)
    gap_pill = Inches(0.2)

    pills_data = [
        ("Step 1 | Crawl & Hash", "Async crawl + SHA-256 fingerprint deduplication"),
        ("Step 2 | OCR Routing", "rho < 120 chars/page triggers Tesseract OCR"),
        ("Step 3 | Dual Index & Serve", "BM25 + 384-d vectors served under 150ms")
    ]

    for i, (p_t, p_d) in enumerate(pills_data):
        xp = Inches(1.05) + i * (w_pill + gap_pill)
        pill = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, xp, y_pills, w_pill, h_pill)
        pill.fill.solid()
        pill.fill.fore_color.rgb = ACCENT_BG
        pill.line.color.rgb = BORDER_COLOR
        
        tx_p = slide5.shapes.add_textbox(xp + Inches(0.1), y_pills + Inches(0.06), w_pill - Inches(0.2), h_pill - Inches(0.12))
        tf_p = tx_p.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = tf_p.margin_top = tf_p.margin_right = tf_p.margin_bottom = 0
        p_pt = tf_p.paragraphs[0]
        p_pt.text = p_t
        p_pt.font.name = FONT_FAMILY
        p_pt.font.size = Pt(8.5)
        p_pt.font.bold = True
        p_pt.font.color.rgb = ACCENT_DARK

        p_pd = tf_p.add_paragraph()
        p_pd.text = p_d
        p_pd.font.name = FONT_FAMILY
        p_pd.font.size = Pt(7.5)
        p_pd.font.color.rgb = TEXT_BODY

    # Right Container Card: Pipeline Execution Breakdown
    w_card_5r = Inches(3.38)
    x_card_5r = Inches(9.15)
    add_card(slide5, x_card_5r, y_pos_5, w_card_5r, h_card_5, "Pipeline Execution Stages", "Five-Stage Ingestion to Decision Lifecycle")

    tx_5r = slide5.shapes.add_textbox(x_card_5r + Inches(0.2), y_pos_5 + Inches(0.65), w_card_5r - Inches(0.4), h_card_5 - Inches(0.75))
    tf_5r = tx_5r.text_frame
    tf_5r.word_wrap = True
    tf_5r.margin_left = tf_5r.margin_top = tf_5r.margin_right = tf_5r.margin_bottom = 0

    panel_guide = [
        ("1. Notice Ingestion & Hash", "Asynchronous crawlers poll portals. SHA-256 fingerprinting skips re-indexing unmodified notices, saving 80% network overhead."),
        ("2. Adaptive OCR Routing", "Computes text density rho. If rho < 120 chars/page, routes to Tesseract OCR; else parses natively via PyPDF in milliseconds."),
        ("3. Entity Normalisation", "Parses Indian monetary scales ('Lakhs', 'Crores', 'p.a.') and deadline timestamps into strict validated schemas."),
        ("4. Dual-Index Persistence", "Stores validated records in SQLite WAL mode. Constructs BM25Okapi inverted index and 384-d semantic vectors."),
        ("5. Serving & Decision Gate", "User profile is screened via deterministic rule checks + RRF rank fusion (k=60), outputting deep links, Hinglish guide & .ics.")
    ]

    for i, (g_title, g_desc) in enumerate(panel_guide):
        p = tf_5r.paragraphs[0] if i == 0 else tf_5r.add_paragraph()
        if i > 0:
            p.space_before = Pt(6)
        p.text = f"• {g_title}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE if i in [1, 4] else TEXT_MAIN

        pd = tf_5r.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {g_desc}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(8.5)
        pd.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 6: HYBRID RETRIEVAL & RANKING ENGINE
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide6, 6)
    add_header(slide6, "Information Retrieval", "Hybrid Search: BM25Okapi Lexical + Dense 384-d Embedding",
               "Reciprocal Rank Fusion unifies sparse term precision with dense conceptual understanding.")

    w_half_6 = Inches(5.7)
    gap_6 = Inches(0.33)
    y_pos_6 = Inches(1.9)
    h_6 = Inches(3.2)

    add_card(slide6, Inches(0.8), y_pos_6, w_half_6, h_6, "Dual Retrieval Strategy", "Sparse Lexical vs. Dense Semantic Models")
    tx_6l = slide6.shapes.add_textbox(Inches(1.05), y_pos_6 + Inches(0.72), w_half_6 - Inches(0.5), h_6 - Inches(0.8))
    tf_6l = tx_6l.text_frame
    tf_6l.word_wrap = True
    tf_6l.margin_left = tf_6l.margin_top = tf_6l.margin_right = tf_6l.margin_bottom = 0

    points_6l = [
        ("Lexical Retrieval (BM25Okapi)", "Matches exact scheme acronyms, quota categories (e.g., 'OBC-NCL', 'PMSSS', 'AICTE Pragati'), and state names. Prevents semantic model false positives on strict administrative acronyms."),
        ("Semantic Retrieval (Dense 384-d Vectors)", "Encodes queries and schemes into 384-dimensional vector space using all-MiniLM-L6-v2. Captures student intent such as 'aid for economically weak female engineering students'."),
        ("The Limitation of Solo Models", "BM25 fails on natural conversational phrasing; Semantic search often misses strict administrative acronyms. Dual hybrid retrieval is required.")
    ]
    for i, (t, d) in enumerate(points_6l):
        p = tf_6l.paragraphs[0] if i == 0 else tf_6l.add_paragraph()
        if i > 0:
            p.space_before = Pt(7)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        pd = tf_6l.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_BODY

    add_card(slide6, Inches(0.8) + w_half_6 + gap_6, y_pos_6, w_half_6, h_6, "Reciprocal Rank Fusion (RRF)", "Non-parametric score normalization")
    tx_6r = slide6.shapes.add_textbox(Inches(0.8) + w_half_6 + gap_6 + Inches(0.25), y_pos_6 + Inches(0.72), w_half_6 - Inches(0.5), h_6 - Inches(0.8))
    tf_6r = tx_6r.text_frame
    tf_6r.word_wrap = True
    tf_6r.margin_left = tf_6r.margin_top = tf_6r.margin_right = tf_6r.margin_bottom = 0

    points_6r = [
        ("RRF Ranking Formula", "RRF(d) = Σ [ 1 / (k + r_m(d)) ]  for each retriever m ∈ {BM25, Dense}, with k = 60."),
        ("Scale Independence", "Eliminates the challenge of calibrating unbounded BM25 scores against cosine similarity intervals [-1, 1]."),
        ("Outlier Robustness", "Ensures that a document ranked moderately well by both retrievers outranks a document that is ranked top in only one."),
        ("Empirical Validation", "Yields higher Mean Reciprocal Rank (MRR@10) and NDCG@10 on benchmark student test queries.")
    ]
    for i, (t, d) in enumerate(points_6r):
        p = tf_6r.paragraphs[0] if i == 0 else tf_6r.add_paragraph()
        if i > 0:
            p.space_before = Pt(6)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        pd = tf_6r.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = ACCENT_BLUE if i == 0 else TEXT_BODY

    bot_card_6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.35))
    bot_card_6.fill.solid()
    bot_card_6.fill.fore_color.rgb = ACCENT_BG
    bot_card_6.line.color.rgb = BORDER_COLOR

    tx_bot_6 = slide6.shapes.add_textbox(Inches(1.05), Inches(5.45), Inches(11.2), Inches(1.15))
    tf_bot_6 = tx_bot_6.text_frame
    tf_bot_6.word_wrap = True
    tf_bot_6.margin_left = tf_bot_6.margin_top = tf_bot_6.margin_right = tf_bot_6.margin_bottom = 0
    pb1 = tf_bot_6.paragraphs[0]
    pb1.text = "HYBRID QUERY LIFECYCLE WALKTHROUGH"
    pb1.font.name = FONT_FAMILY
    pb1.font.size = Pt(10)
    pb1.font.bold = True
    pb1.font.color.rgb = ACCENT_DARK

    pb2 = tf_bot_6.add_paragraph()
    pb2.space_before = Pt(3)
    pb2.text = "1. Student Query ➔ 2. Parallel Dispatch (BM25Okapi sparse search + all-MiniLM-L6-v2 vector dot product) ➔ 3. RRF Rank Fusion (k=60) ➔ 4. Deterministic Eligibility Gate (Drops unqualified schemes) ➔ 5. Top Ranked Opportunities Returned (< 150ms total response time)."
    pb2.font.name = FONT_FAMILY
    pb2.font.size = Pt(10)
    pb2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 7: ADAPTIVE OCR & DETERMINISTIC ELIGIBILITY ENGINE
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide7, 7)
    add_header(slide7, "Verification & Rules", "Adaptive OCR Routing & Deterministic Eligibility Engine",
               "Zero-hallucination policy: Density-driven text extraction combined with strict mathematical eligibility gates.")

    w_7 = Inches(5.7)
    gap_7 = Inches(0.33)
    y_pos_7 = Inches(1.9)
    h_7 = Inches(4.75)

    add_card(slide7, Inches(0.8), y_pos_7, w_7, h_7, "Adaptive Document Ingestion", "Density-based OCR routing for official notices")
    tx_7l = slide7.shapes.add_textbox(Inches(1.05), y_pos_7 + Inches(0.72), w_7 - Inches(0.5), h_7 - Inches(0.8))
    tf_7l = tx_7l.text_frame
    tf_7l.word_wrap = True
    tf_7l.margin_left = tf_7l.margin_top = tf_7l.margin_right = tf_7l.margin_bottom = 0

    points_7l = [
        ("The Document Challenge", "Government gazettes arrive in two formats: modern digital PDFs and degraded scanned paper circulars with official stamps and seals."),
        ("Density Metric rho < 120 chars/page", "Calculates character density per page. If rho >= 120, native PyPDF/pdfplumber extracts pristine electronic text in milliseconds."),
        ("Tesseract OCR Fallback", "If rho < 120 chars/page (scanned document), pages are rendered to high-res TIFF images and routed through Tesseract OCR."),
        ("Efficiency Advantage", "Bypassing OCR for 70%+ of digital notices eliminates processing bottlenecks and minimizes server compute overhead drastically."),
        ("Entity Normalisation", "Parses Indian financial terminology ('Lakhs', 'Crores', 'p.a.') and cutoff dates into structured database schemas.")
    ]
    for i, (t, d) in enumerate(points_7l):
        p = tf_7l.paragraphs[0] if i == 0 else tf_7l.add_paragraph()
        if i > 0:
            p.space_before = Pt(8)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        pd = tf_7l.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_BODY

    add_card(slide7, Inches(0.8) + w_7 + gap_7, y_pos_7, w_7, h_7, "Strict Eligibility Verification Engine", "Guaranteed zero hallucination in qualification")
    tx_7r = slide7.shapes.add_textbox(Inches(0.8) + w_7 + gap_7 + Inches(0.25), y_pos_7 + Inches(0.72), w_7 - Inches(0.5), h_7 - Inches(0.8))
    tf_7r = tx_7r.text_frame
    tf_7r.word_wrap = True
    tf_7r.margin_left = tf_7r.margin_top = tf_7r.margin_right = tf_7r.margin_bottom = 0

    points_7r = [
        ("Hard Gate vs. Probabilistic Gate", "LLMs frequently hallucinate eligibility requirements. MadadgaarAI isolates all eligibility logic into deterministic code."),
        ("12th Marks Aggregate Check", "Evaluates: student_percentage >= scheme_minimum_cutoff (e.g. 85.0% aggregate requirement)."),
        ("Family Annual Income Cap", "Evaluates: student_family_income <= scheme_income_cap (e.g. ₹6,00,000 / year threshold)."),
        ("Demographic & State Quotas", "Evaluates strict categorical constraints: Domicile state match, gender cohorts (Girls-only schemes), and caste quotas."),
        ("Candidate Filtering", "Any scheme where conditions evaluate to FALSE is immediately dropped from the final recommendation feed.")
    ]
    for i, (t, d) in enumerate(points_7r):
        p = tf_7r.paragraphs[0] if i == 0 else tf_7r.add_paragraph()
        if i > 0:
            p.space_before = Pt(8)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        pd = tf_7r.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = SUCCESS_TEXT if "Hard Gate" in t else TEXT_BODY

    # ==========================================
    # SLIDE 8: MULTILINGUAL ADVISORY & VALUE-ADDED SERVICES
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide8, 8)
    add_header(slide8, "Student Assistance", "Empowering Students: Saral Guide, Vidyarthi AI & Calendar Export",
               "Translating bureaucratic legalese into clear conversational Hinglish with actionable checklists and deadline alerts.")

    col_w_8 = Inches(3.68)
    gap_8 = Inches(0.34)
    y_pos_8 = Inches(1.9)
    card_h_8 = Inches(4.75)

    assist_cards = [
        ("Feature 01", "Saral Guide (Hinglish)", [
             ("Jargon Translation", "Deconstructs 20-page legal gazette circulars into simple, everyday Hinglish bullet points."),
             ("Cognitive Relief", "Tier-2/3 rural students can understand guidelines in 30 seconds without consulting third-party agents."),
             ("Key Eligibility Summary", "Highlights cutoffs, annual grant amounts, and exact renewal clauses in straightforward conversational prose."),
             ("Zero Ambiguity", "Explains complex bond agreements and conditions in transparent vernacular language.")
         ]),
        ("Feature 02", "Dynamic Document Checklist", [
             ("Automated Requirement Mapping", "Dynamically generates exact list of documents required for the chosen scheme."),
             ("Issuing Authority Details", "Specifies valid issuing offices (e.g. Income Certificate from Tehsildar/SDM, not self-affidavit)."),
             ("Document Readiness Status", "Provides interactive checkboxes for students to track preparation."),
             ("Zero Rejection Rate", "Prevents application rejection caused by missing ancillary certificates.")
         ]),
        ("Feature 03", "Vidyarthi AI & RFC 5545", [
             ("Autonomous Guidance Agent", "An intelligent companion addressing students' context-specific doubts and queries."),
             ("Alternative Scheme Discovery", "If a student fails an eligibility condition, Vidyarthi AI suggests alternative open schemes."),
             ("RFC 5545 Calendar Export", "Generates downloadable .ics calendar files syncing scheme deadlines with Google/Apple Calendar."),
             ("Accessible Web Interface", "Integrated directly within the Neo-Brutalist dashboard modal for single-click access.")
         ])
    ]

    for i, (tag, head, bullets) in enumerate(assist_cards):
        x = Inches(0.8) + i * (col_w_8 + gap_8)
        add_card(slide8, x, y_pos_8, col_w_8, card_h_8)

        tx = slide8.shapes.add_textbox(x + Inches(0.28), y_pos_8 + Inches(0.25), col_w_8 - Inches(0.56), card_h_8 - Inches(0.5))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = tag.upper()
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p_h = tf.add_paragraph()
        p_h.space_before = Pt(4)
        p_h.text = head
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_MAIN

        for b_title, b_desc in bullets:
            pb_t = tf.add_paragraph()
            pb_t.space_before = Pt(10)
            pb_t.text = f"• {b_title}"
            pb_t.font.name = FONT_FAMILY
            pb_t.font.size = Pt(10.5)
            pb_t.font.bold = True
            pb_t.font.color.rgb = TEXT_MAIN

            pb_d = tf.add_paragraph()
            pb_d.space_before = Pt(1)
            pb_d.text = f"  {b_desc}"
            pb_d.font.name = FONT_FAMILY
            pb_d.font.size = Pt(9.5)
            pb_d.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 9: TECH STACK & EXPERIMENTAL EVALUATION
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide9, 9)
    add_header(slide9, "Tech Stack & Evaluation", "Production Technology Stack & Experimental Performance",
               "Validation across asynchronous retrieval speed, deterministic accuracy, and lightweight system footprint.")

    y_stat_9 = Inches(1.9)
    h_stat_9 = Inches(1.25)
    w_stat_9 = Inches(3.68)
    gap_stat_9 = Inches(0.34)

    stat_boxes_9 = [
        ("< 150 ms", "Average Query Response Latency", "Parallel execution of sparse BM25Okapi and dense vector dot product."),
        ("0% False Positives", "Deterministic Verification Accuracy", "Mathematical filtering eliminates unqualified scheme promises."),
        ("70% Less Compute", "OCR Optimization via Density Checks", "Native PyPDF avoids unnecessary OCR on clean electronic circulars.")
    ]

    for i, (s_num, s_title, s_sub) in enumerate(stat_boxes_9):
        x = Inches(0.8) + i * (w_stat_9 + gap_stat_9)
        c = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_stat_9, w_stat_9, h_stat_9)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR

        tx = slide9.shapes.add_textbox(x + Inches(0.2), y_stat_9 + Inches(0.12), w_stat_9 - Inches(0.4), h_stat_9 - Inches(0.24))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = s_num
        p.font.name = FONT_FAMILY
        p.font.size = Pt(21)
        p.font.bold = True
        p.font.color.rgb = ACCENT_DARK

        p_t = tf.add_paragraph()
        p_t.space_before = Pt(2)
        p_t.text = s_title
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_s = tf.add_paragraph()
        p_s.space_before = Pt(1)
        p_s.text = s_sub
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(8.5)
        p_s.font.color.rgb = TEXT_MUTED

    y_bot_9 = Inches(3.35)
    h_bot_9 = Inches(3.3)
    w_half_9 = Inches(5.7)
    gap_half_9 = Inches(0.33)

    add_card(slide9, Inches(0.8), y_bot_9, w_half_9, h_bot_9, "Production Technology Stack", "Modular, decoupled open-source architecture")
    tx_9l = slide9.shapes.add_textbox(Inches(1.05), y_bot_9 + Inches(0.7), w_half_9 - Inches(0.5), h_bot_9 - Inches(0.8))
    tf_9l = tx_9l.text_frame
    tf_9l.word_wrap = True
    tf_9l.margin_left = tf_9l.margin_top = tf_9l.margin_right = tf_9l.margin_bottom = 0

    stack_rows = [
        ("Backend Framework", "Python 3.12+, FastAPI, Uvicorn (High concurrency async engine)"),
        ("Validation & Schema", "Pydantic v2 (Strict data models, ENUM typing, 0% unhandled 422s)"),
        ("Storage Engine", "SQLite (WAL Mode) + JSON/CSV (Zero DevOps overhead, ACID compliance)"),
        ("Dual-Layer Search", "BM25Okapi + SentenceTransformers (all-MiniLM-L6-v2 384-d vectors)"),
        ("Document Parsing", "PyPDF/pdfplumber + Tesseract OCR (Density threshold rho < 120)"),
        ("Frontend & UX", "Neo-Brutalist Dashboard (Vanilla CSS & JS, zero heavy bundle bloat)")
    ]
    for i, (cat, val) in enumerate(stack_rows):
        p = tf_9l.paragraphs[0] if i == 0 else tf_9l.add_paragraph()
        if i > 0:
            p.space_before = Pt(5)
        p.text = f"• {cat}: "
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        r = p.add_run()
        r.text = val
        r.font.name = FONT_FAMILY
        r.font.size = Pt(9)
        r.font.bold = False
        r.font.color.rgb = TEXT_BODY

    add_card(slide9, Inches(0.8) + w_half_9 + gap_half_9, y_bot_9, w_half_9, h_bot_9, "Experimental Validation Highlights", "Benchmarked across accuracy and real test cases")
    tx_9r = slide9.shapes.add_textbox(Inches(0.8) + w_half_9 + gap_half_9 + Inches(0.25), y_bot_9 + Inches(0.7), w_half_9 - Inches(0.5), h_bot_9 - Inches(0.8))
    tf_9r = tx_9r.text_frame
    tf_9r.word_wrap = True
    tf_9r.margin_left = tf_9r.margin_top = tf_9r.margin_right = tf_9r.margin_bottom = 0

    eval_points = [
        ("Hybrid RRF vs Single Models", "BM25 alone struggled with vague queries; Dense vectors alone suffered from state-level category drift. RRF achieved superior MRR@10 and NDCG@10."),
        ("OCR Throughput Gains", "By applying density checks (rho < 120), 73% of digital PDF circulars bypassed OCR entirely, achieving 3.8x faster ingestion speed."),
        ("Live Prototype Deployment", "Full end-to-end integration verified: match API, Saral guide modal, dynamic checklist generation, and student profile scoring.")
    ]
    for i, (t, d) in enumerate(eval_points):
        p = tf_9r.paragraphs[0] if i == 0 else tf_9r.add_paragraph()
        if i > 0:
            p.space_before = Pt(7)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE if i == 0 else TEXT_MAIN

        pd = tf_9r.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 10: FUTURE WORK & CONCLUSION
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    apply_slide_base(slide10, 10)
    add_header(slide10, "Roadmap & Conclusion", "Phase 2 Development Roadmap & Final Conclusion",
               "Strategic extensions planned for the final major project submission and thesis defense.")

    w_10 = Inches(5.7)
    gap_10 = Inches(0.33)
    y_pos_10 = Inches(1.9)
    h_10 = Inches(4.75)

    add_card(slide10, Inches(0.8), y_pos_10, w_10, h_10, "Future Work for Final Submission", "Four key technical modules slated for Phase 2")
    tx_10l = slide10.shapes.add_textbox(Inches(1.05), y_pos_10 + Inches(0.72), w_10 - Inches(0.5), h_10 - Inches(0.8))
    tf_10l = tx_10l.text_frame
    tf_10l.word_wrap = True
    tf_10l.margin_left = tf_10l.margin_top = tf_10l.margin_right = tf_10l.margin_bottom = 0

    future_points = [
        ("1. Autonomous Form-Filling Agents", "Develop an orchestration framework with browser automation (Playwright/Selenium) to assist students in pre-populating portal application forms."),
        ("2. Document Authenticity Verification", "Integrate a computer-vision tampering detection model to verify official stamps, QR signatures, and dates on uploaded certificates."),
        ("3. WhatsApp & SMS Chatbot Outreach", "Deploy low-bandwidth conversational access channels to serve rural students in regions with limited broadband access."),
        ("4. Vernacular Voice Query Engine", "Incorporate lightweight open-source Indic Whisper models for hands-free regional voice queries in Hindi, Bengali, Tamil, and Marathi.")
    ]
    for i, (t, d) in enumerate(future_points):
        p = tf_10l.paragraphs[0] if i == 0 else tf_10l.add_paragraph()
        if i > 0:
            p.space_before = Pt(8)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        pd = tf_10l.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_BODY

    add_card(slide10, Inches(0.8) + w_10 + gap_10, y_pos_10, w_10, h_10, "Major Project Summary & Impact", "Bridging India's educational opportunity divide")
    tx_10r = slide10.shapes.add_textbox(Inches(0.8) + w_10 + gap_10 + Inches(0.25), y_pos_10 + Inches(0.72), w_10 - Inches(0.5), h_10 - Inches(0.8))
    tf_10r = tx_10r.text_frame
    tf_10r.word_wrap = True
    tf_10r.margin_left = tf_10r.margin_top = tf_10r.margin_right = tf_10r.margin_bottom = 0

    conc_points = [
        ("Bridging the Welfare Gap", "MadadgaarAI eliminates the complex information asymmetry that prevents underrepresented students from discovering and claiming financial support."),
        ("Engineering Rigor", "By combining deterministic verification (zero hallucination) with state-of-the-art hybrid neural search (BM25Okapi + all-MiniLM-L6-v2), the platform guarantees speed and reliability."),
        ("Equitable Access", "Saral Guide demystifies bureaucratic barriers into actionable Hinglish guidance, making government funding truly accessible."),
        ("Demo-Ready Execution", "A fully working, live-tested prototype is functional with verified API contracts, clean UI, and real-time retrieval performance.")
    ]
    for i, (t, d) in enumerate(conc_points):
        p = tf_10r.paragraphs[0] if i == 0 else tf_10r.add_paragraph()
        if i > 0:
            p.space_before = Pt(8)
        p.text = f"• {t}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        pd = tf_10r.add_paragraph()
        pd.space_before = Pt(1)
        pd.text = f"  {d}"
        pd.font.name = FONT_FAMILY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_BODY

    # Save presentation
    output_pptx = "MadadgaarAI_Presentation.pptx"
    prs.save(output_pptx)
    print(f"Presentation saved successfully to {output_pptx}")

if __name__ == "__main__":
    create_deck()
